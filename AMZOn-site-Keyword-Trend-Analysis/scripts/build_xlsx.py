#!/usr/bin/env python3
"""Build the trend-selection Excel report from payload.json (structure: references/output-schema.md).

Usage: python build_xlsx.py <payload.json> [<output.xlsx>]
Default output: <meta.out_dir>/{topic}_趋势选品_v{version}_{YYYYMMDD}_{AM|PM}{hhmm}.xlsx
Sheets: 说明 / 候选总览 / 候选卡 / 预测 / 趋势序列 / 证据台账 / 红旗 / 监测清单
"""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
VERDICT_FILL = {"推荐": "E2F0D9", "观察": "FFF3CD", "放弃": "F8D7DA"}
STAMP_RE = re.compile(r"_\d{8}_(AM|PM)\d{4}$")
TAIL = "本结论不含供给与利润判断。"


def stamp(now=None):
    now = now or datetime.now()
    return f"{now:%Y%m%d}_{'AM' if now.hour < 12 else 'PM'}{now.hour % 12 or 12:02d}{now:%M}"


def style_table(ws, widths):
    for row in ws.iter_rows():
        for c in row:
            c.font = BODY
            c.alignment = WRAP
            c.border = BORDER
    for c in ws[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def join(v):
    return "\n".join(f"· {x}" for x in v) if isinstance(v, list) else (v or "")


def main(src, out):
    data = json.load(open(src, encoding="utf-8"))
    meta = data.get("meta", {})
    cands = data.get("candidates", [])
    topic = meta.get("topic", "趋势选品")
    safe = re.sub(r'[\\/:*?"<>|]', "_", topic)
    if not out:
        out = os.path.join(meta.get("out_dir", "."), f"{safe if '趋势选品' in safe else safe + '_趋势选品'}_v{meta.get('version', 1)}.xlsx")
    root, ext = os.path.splitext(out)
    if not STAMP_RE.search(root):
        out = f"{root}_{stamp()}{ext or '.xlsx'}"
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    n = {v: sum(1 for c in cands if c.get("verdict") == v) for v in VERDICT_FILL}
    rows = [
        ("主题", topic),
        ("站点", meta.get("marketplace", "")),
        ("数据截至", meta.get("as_of", "")),
        ("预测时点", f"{meta.get('forecast_target_date', '')}（+{meta.get('forecast_horizon_months', '')} 个月）"),
        ("入口", meta.get("entry", "")),
        ("筛选阈值", meta.get("thresholds", "")),
        ("候选数", f"推荐 {n['推荐']}　观察 {n['观察']}　放弃 {n['放弃']}"),
        ("使用的工具", "、".join(meta.get("tools_used", []))),
        ("数据缺口/降级", join(meta.get("unknowns", []))),
        ("范围声明", "只看需求端趋势：不含成本、供应链、库存、广告盈亏。" + TAIL),
        ("证据分级", "FACT 原始返回 / ESTIMATE 第三方估算 / ASSUMPTION 分析假设 / UNKNOWN 未获得。预测为定性判断，不是承诺。"),
        ("红旗", "功效方向与品牌词红旗只作提示，不影响趋势结论。"),
    ]
    ws.append(["项目", "内容"])
    for r in rows:
        ws.append(list(r))
    style_table(ws, [16, 100])

    ws = wb.create_sheet("候选总览")
    ws.append(["编号", "关键词", "中文", "原型", "预测方向", "阶段", "置信度", "结论", "理由", "红旗数", "UNKNOWN"])
    for c in cands:
        f = c.get("forecast", {})
        ws.append([c.get("id"), c.get("keyword"), c.get("zh"), c.get("archetype"), f.get("direction"), f.get("stage"),
                   f.get("confidence"), c.get("verdict"), c.get("verdict_reason"), len(c.get("flags", [])),
                   join(c.get("unknowns", []))])
    style_table(ws, [7, 24, 16, 14, 10, 12, 8, 8, 50, 7, 30])
    for row in ws.iter_rows(min_row=2):
        fill = VERDICT_FILL.get(row[7].value)
        if fill:
            row[7].fill = PatternFill("solid", fgColor=fill)

    ws = wb.create_sheet("候选卡")
    ws.append(["编号", "关键词", "入口", "需求真实性", "趋势诊断", "进入窗口（需求流向）", "需求缺口"])
    for c in cands:
        ws.append([c.get("id"), c.get("keyword"), c.get("entry"), c.get("demand"), c.get("trend"),
                   c.get("window"), c.get("gap")])
    style_table(ws, [7, 24, 8, 40, 50, 40, 40])

    ws = wb.create_sheet("预测")
    ws.append(["编号", "关键词", "方向", "阶段", "预测时点位置", "距峰（周）", "峰值依据", "保守", "基准", "乐观",
               "置信度", "置信依据", "证实信号", "证伪信号", "备注"])
    for c in cands:
        f = c.get("forecast", {})
        s = f.get("scenarios", {})
        ws.append([c.get("id"), c.get("keyword"), f.get("direction"), f.get("stage"), f.get("target_position"),
                   f.get("weeks_to_peak"), f.get("peak_basis"), s.get("conservative"), s.get("base"),
                   s.get("optimistic"), f.get("confidence"), f.get("confidence_basis"), join(f.get("confirm", [])),
                   join(f.get("refute", [])), TAIL])
    style_table(ws, [7, 22, 8, 12, 30, 9, 26, 30, 30, 30, 8, 26, 34, 34, 22])

    ws = wb.create_sheet("趋势序列")
    ws.append(["编号", "关键词", "序列", "来源", "粒度", "单位", "口径", "日期", "数值"])
    for c in cands:
        for s in c.get("series", []):
            for d, v in zip(s.get("dates", []), s.get("values", [])):
                ws.append([c.get("id"), c.get("keyword"), s.get("name"), s.get("source"), s.get("granularity"),
                           s.get("unit"), s.get("tag"), d, v])
    style_table(ws, [7, 22, 22, 34, 8, 10, 10, 12, 14])

    ws = wb.create_sheet("证据台账")
    ws.append(["编号", "工具", "参数", "数据窗口", "口径", "发现"])
    for c in cands:
        for e in c.get("evidence", []):
            ws.append([c.get("id"), e.get("tool"), e.get("params"), e.get("data_period"), e.get("tag"), e.get("finding")])
    style_table(ws, [7, 36, 40, 22, 10, 60])

    ws = wb.create_sheet("红旗")
    ws.append(["编号", "关键词", "类型", "说明"])
    for c in cands:
        for f in c.get("flags", []):
            ws.append([c.get("id"), c.get("keyword"), f.get("type"), f.get("text")])
    style_table(ws, [7, 24, 24, 80])

    ws = wb.create_sheet("监测清单")
    ws.append(["编号", "复查什么", "何时", "触发条件"])
    for w in data.get("watchlist", []):
        ws.append([w.get("candidate"), w.get("what"), w.get("when"), w.get("trigger")])
    style_table(ws, [7, 40, 20, 60])

    wb.save(out)
    print(out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
