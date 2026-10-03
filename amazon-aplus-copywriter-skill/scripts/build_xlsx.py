#!/usr/bin/env python3
"""Build the A+ / Brand Story copy workbook from a JSON file.

Usage: python3 build_xlsx.py copy.json output.xlsx
       python3 build_xlsx.py --stamp-only        # print the timestamp only, e.g. 20261001_PM1030

The output name gets a generation timestamp before ".xlsx" unless it already has one:
  B0XXXX_A+文案_v1.xlsx -> B0XXXX_A+文案_v1_20261001_PM1030.xlsx  (12-hour clock, AM/PM before hhmm)
Name copy.json / Word / Markdown deliverables with the same stamp (get it via --stamp-only).

JSON structure (all text may contain real newlines; they become in-cell line breaks):
{
  "meta": {
    "title": "B0XXXX｜A+ 与 Brand Story 文案 v1",
    "subtitle": "ASIN：...｜品牌 产品 规格｜编写日期：YYYY-MM-DD",
    "decisions": ["品牌名只用 ZAB", "..."],          # 用户在 plan 阶段拍板的事项
    "rules": ["每个数字都写清对象、单位和每份基准", "..."],  # 规则卡摘要
    "notes": ["尺寸提示……", "..."],                  # 其他说明（可选）
    "claims": {"allowed": ["energy metabolism"],     # 本产品功效方向（plan 确定，见 references/ingredient-research.md）
               "blocked": ["gut health", "immune"]}  # scan_copy 按 blocked 拦截越界说法
  },
  "rows": [                                          # 每个文字位一行
    {"id": "A1-01", "group": "A+图1", "file": "A+纯图1.png", "slot": "左侧引语",
     "placeholder": "占位原文（新增位写“（无）”，不适用写“—”）",
     "copy": "新英文文案", "alt": "备选英文", "zh": "中文释义",
     "facts": "F01, F02", "note": "合规/设计备注",
     "voice": "brand"}                               # 可选：品牌调性文字位（含其 ALT 行），scan_copy 会做专项检查
  ],
  "marketing": [                                     # 可选：营销话术候选池（见 references/persuasion-psychology.md）
    {"id": "MK-01", "principle": "损失厌恶", "mechanism": "人对失去比得到更敏感",
     "channel": "A+|站外", "level": "✅|⚠️|⛔", "copy": "英文话术", "alt": "备选", "zh": "中文释义",
     "placement": "M-2 底部小字 / 【新增】角标", "facts": "F02", "risk": "风险说明",
     "positive_read": "正读（损失厌恶行必填）", "negative_read": "反读：顾客自己推出来，图上不写"}
  ],
  "explainers": [                                    # 可选：成分科普（见 references/ingredient-eli5.md）
    {"id": "EX-01", "question": "What is [成分]?", "essence": "一句话本质", "analogy": "类比",
     "detail": "展开 2-3 句", "so_what": "和我有什么关系", "on_image": "图上标题\n图上正文", "zh": "中文释义",
     "placement": "M-3 / 【新增】教育图", "facts": "F11", "source": "证据类型＋出处方向",
     "level": "✅|⚠️", "risk": "风险说明"}
  ],
  "excipients": {                                    # 可选：辅料与剂型（见 references/dosage-form-excipients.md）
    "form": "硬胶囊（vegetable cellulose 植物胶囊）", "source": "标签图",
    "items": [
      {"id": "EXC-A", "angle": "外壳/基质", "headline": "英文标题", "subheadline": "英文副标题",
       "points": "支撑卖点（换行分隔）", "zh": "中文释义", "facts": "F14", "level": "✅|⚠️", "risk": "风险说明"}
    ]
  },
  "issues": [                                        # 设计修改清单
    {"priority": "必改|建议|确认", "location": "...", "problem": "...", "fix": "..."}
  ],
  "facts": [                                         # 事实依据
    {"id": "F01", "fact": "...", "source": "...", "note": "..."}
  ]
}
Markup convention in copy: **bold**, *accent/italic*, trailing lone * = structure/function claim asterisk.
"""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
PALETTE = ["FDF2F2", "F2F6FC", "F4F9F1", "FBF7EE", "F5F0FA", "EEF8F8", "F3F3F3", "FFF8E6"]
PRI_FILL = {"必改": "F8D7DA", "建议": "FFF3CD", "确认": "D9E7F7"}
LEVEL_FILL = {"✅": "D4EDDA", "⚠️": "FFF3CD", "⛔": "F8D7DA"}
APLUS_REDLINES = [
    "A+ 红线摘要（以卖家后台当前 A+ 内容指南为准）：以下内容放进 A+ 图片会被拒审——",
    "· 价格/促销：$、price、free、bonus、affordable、discount、sale",
    "· 购买引导：buy now、add to cart、get yours、order now",
    "· 时效：now、new、latest、limited、hurry、only N left、today only、节日",
    "· 评价/排名：reviews、stars、best-selling、#1、top-rated",
    "· 保证/退款：guarantee、refund、money-back、risk-free",
    "⛔ 行只能用于广告、社媒、站外落地页；方括号 [N] 处必须填入真实数据后才能用。",
]


STAMP_RE = re.compile(r"_\d{8}_(AM|PM)\d{4}$")


def stamp(now=None):
    """Generation timestamp like 20261001_PM1030 (12-hour clock; locale-independent)."""
    now = now or datetime.now()
    hour12 = now.hour % 12 or 12
    return f"{now:%Y%m%d}_{'AM' if now.hour < 12 else 'PM'}{hour12:02d}{now:%M}"


def stamped(path):
    root, ext = os.path.splitext(path)
    return path if STAMP_RE.search(root) else f"{root}_{stamp()}{ext or '.xlsx'}"


def style_table(ws, widths):
    for row in ws.iter_rows():
        for c in row:
            c.font = BODY
            c.alignment = WRAP
            c.border = BORDER
    for c in ws[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def main(src, out):
    out = stamped(out)
    data = json.load(open(src, encoding="utf-8"))
    meta = data.get("meta", {})
    wb = Workbook()

    # --- 说明 ---
    ws = wb.active
    ws.title = "说明"
    lines = [(meta.get("title", "A+ 文案"), "title"), (meta.get("subtitle", ""), None), ("", None),
             ("怎么用这份表", "h")]
    lines += [(t, None) for t in [
        "1. 「文案总表」每行对应构图里的一个文字位。设计直接复制E列（新英文文案），F列是备选句。",
        "2. 标注约定：**…** 表示加粗，*…* 表示强调色/斜体，句尾单独的 * 是功能声明星号（要和FDA免责声明对应）。单元格里怎么换行，排版就怎么换行。",
        "3. H列/I列用 LEN 公式统计新文案和占位原文的字符数，J列是差值（正数=比原稿长）。差值超过 +10 的位置，需要设计检查是否放得下。",
        "4. 「设计修改清单」列出图片或素材层面要改的地方（不是文案），上线前需要逐项关掉。",
        "5. 「事实依据」列出每条文案引用的事实编号（F01…）及出处。",
    ]]
    if data.get("marketing"):
        lines.append(("6. 「营销话术」是按消费心理学原理写的候选池，不替换主文案。每条注明建议位置和合规等级："
                      "✅ 可直接用、⚠️ 需确认（如年龄框架，需 * 和 FDA 声明）、⛔ 不可放 A+（只用于站外）。"
                      "每张图最多选 1 条。损失厌恶行另有正读/反读两列：反读是顾客自己推出来的意思，不写在图上。", None))
    if data.get("explainers"):
        lines.append(("7. 「成分科普」按顾客的提问顺序，用 ELI5 的结构（一句话本质 → 类比 → 展开 → 和我有什么关系）"
                      "讲成分如何作用。G 列「图上短版」可直接用于 A+ 教育图（第一行为标题），其余列可用于 FAQ、详情或客服回答。"
                      "研究出处只在 L 列，供核对，不上图。", None))
    if (data.get("excipients") or {}).get("items"):
        lines.append(("8. 「辅料与剂型」按标签上的剂型和 Other Ingredients，从三个角度（外壳/基质、辅料清单、各司其职）"
                      "各写一版 Headline + Subheadline + 支撑卖点，可用于 Supplement Facts 配套图或成分透明图。"
                      "剂型变了（软胶囊、软糖、片剂…）这张表会随之重写。", None))
    claims = meta.get("claims") or {}
    if claims:
        lines += [("", None), ("本产品功效方向", "h"),
                  ("· 允许（allowed）：" + "、".join(claims.get("allowed") or []), None),
                  ("· 不写（blocked）：" + "、".join(claims.get("blocked") or []), None)]
    if meta.get("decisions"):
        lines += [("", None), ("本次已确认的决定", "h")] + [("· " + d, None) for d in meta["decisions"]]
    if meta.get("rules"):
        lines += [("", None), ("写作规则（规则卡摘要）", "h")] + [("· " + r, None) for r in meta["rules"]]
    if meta.get("notes"):
        lines += [("", None), ("其他说明", "h")] + [("· " + n, None) for n in meta["notes"]]
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=kind is not None, size=14 if kind == "title" else (11 if kind else 10))
        c.alignment = WRAP
    ws.column_dimensions["A"].width = 120

    # --- 文案总表 ---
    ws = wb.create_sheet("文案总表")
    headers = ["编号", "图片", "构图文件", "文字位置", "新英文文案", "备选英文", "中文释义",
               "新文案字符数", "占位原文字符数", "差值", "占位原文（设计稿）", "事实依据", "合规 / 设计备注"]
    ws.append(headers)
    groups = []
    for r in data["rows"]:
        g = r.get("group", "")
        if g not in groups:
            groups.append(g)
        ws.append([r.get("id", ""), g, r.get("file", ""), r.get("slot", ""), r.get("copy", ""),
                   r.get("alt", ""), r.get("zh", ""), None, None, None, r.get("placeholder", ""),
                   r.get("facts", ""), r.get("note", "")])
        n = ws.max_row
        ws.cell(n, 8).value = f"=LEN(E{n})"
        ws.cell(n, 9).value = f'=IF(OR(K{n}="—",K{n}="",LEFT(K{n},1)="（"),"",LEN(K{n}))'
        ws.cell(n, 10).value = f'=IF(I{n}="","",H{n}-I{n})'
    style_table(ws, [8, 10, 16, 24, 44, 34, 28, 9, 9, 7, 34, 12, 52])
    for n in range(2, ws.max_row + 1):
        fill = PatternFill("solid", fgColor=PALETTE[groups.index(ws.cell(n, 2).value) % len(PALETTE)])
        for col in range(1, len(headers) + 1):
            ws.cell(n, col).fill = fill
        ws.cell(n, 5).font = Font(name=FONT, bold=True, size=10, color="C00000")
        ws.cell(n, 11).font = Font(name=FONT, size=9, color="808080")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"
    ws.row_dimensions[1].height = 30

    # --- 营销话术（可选）---
    marketing = data.get("marketing") or []
    if marketing:
        ws = wb.create_sheet("营销话术")
        mk_headers = ["编号", "心理学原理", "机制一句话", "渠道", "合规等级", "英文话术", "备选英文", "中文释义",
                      "正读（顾客读到的）", "反读（顾客自己推，图上不写）", "建议位置", "字符数", "事实依据", "风险说明"]
        ws.append(mk_headers)
        ordered = ([m for m in marketing if m.get("channel") != "站外"] +
                   [m for m in marketing if m.get("channel") == "站外"])
        principles = []
        for m in ordered:
            p = m.get("principle", "")
            if p not in principles:
                principles.append(p)
            ws.append([m.get("id", ""), p, m.get("mechanism", ""), m.get("channel", "A+"), m.get("level", ""),
                       m.get("copy", ""), m.get("alt", ""), m.get("zh", ""), m.get("positive_read", ""),
                       m.get("negative_read", ""), m.get("placement", ""), None, m.get("facts", ""), m.get("risk", "")])
            ws.cell(ws.max_row, 12).value = f"=LEN(F{ws.max_row})"
        style_table(ws, [8, 12, 26, 7, 9, 44, 34, 28, 28, 30, 24, 8, 10, 40])
        for n in range(2, ws.max_row + 1):
            fill = PatternFill("solid", fgColor=PALETTE[principles.index(ws.cell(n, 2).value) % len(PALETTE)])
            for col in range(1, len(mk_headers) + 1):
                ws.cell(n, col).fill = fill
            lvl = str(ws.cell(n, 5).value or "")
            for mark, color in LEVEL_FILL.items():
                if lvl.startswith(mark):
                    ws.cell(n, 5).fill = PatternFill("solid", fgColor=color)
            ws.cell(n, 6).font = Font(name=FONT, bold=True, size=10, color="C00000")
        ws.freeze_panes = "F2"
        ws.auto_filter.ref = f"A1:N{ws.max_row}"
        ws.row_dimensions[1].height = 30
        start = ws.max_row + 2
        for i, t in enumerate(APLUS_REDLINES):
            c = ws.cell(start + i, 1, value=t)
            c.font = Font(name=FONT, bold=i == 0, size=10)

    # --- 成分科普（可选）---
    explainers = data.get("explainers") or []
    if explainers:
        ws = wb.create_sheet("成分科普")
        ex_headers = ["编号", "顾客问题", "一句话本质", "类比", "展开", "So what（和我有什么关系）", "图上短版",
                      "中文释义", "建议位置", "图上短版字符数", "事实依据", "证据类型 / 出处", "合规等级", "风险说明"]
        ws.append(ex_headers)
        for e in explainers:
            ws.append([e.get("id", ""), e.get("question", ""), e.get("essence", ""), e.get("analogy", ""),
                       e.get("detail", ""), e.get("so_what", ""), e.get("on_image", ""), e.get("zh", ""),
                       e.get("placement", ""), None, e.get("facts", ""), e.get("source", ""),
                       e.get("level", ""), e.get("risk", "")])
            ws.cell(ws.max_row, 10).value = f"=LEN(G{ws.max_row})"
        style_table(ws, [7, 22, 34, 30, 44, 34, 36, 34, 20, 9, 10, 40, 8, 32])
        for n in range(2, ws.max_row + 1):
            fill = PatternFill("solid", fgColor=PALETTE[(n - 2) % len(PALETTE)])
            for col in range(1, len(ex_headers) + 1):
                ws.cell(n, col).fill = fill
            lvl = str(ws.cell(n, 13).value or "")
            for mark, color in LEVEL_FILL.items():
                if lvl.startswith(mark):
                    ws.cell(n, 13).fill = PatternFill("solid", fgColor=color)
            ws.cell(n, 7).font = Font(name=FONT, bold=True, size=10, color="C00000")
            ws.cell(n, 12).font = Font(name=FONT, size=9, color="808080")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:N{ws.max_row}"
        ws.row_dimensions[1].height = 30

    # --- 辅料与剂型（可选）---
    exc = data.get("excipients") or {}
    if exc.get("items"):
        ws = wb.create_sheet("辅料与剂型")
        exc_headers = ["编号", "角度", "Headline", "Subheadline", "支撑卖点", "中文释义", "Headline 字符数",
                       "Subheadline 字符数", "事实依据", "合规等级", "风险说明"]
        ws.append(exc_headers)
        for e in exc["items"]:
            ws.append([e.get("id", ""), e.get("angle", ""), e.get("headline", ""), e.get("subheadline", ""),
                       e.get("points", ""), e.get("zh", ""), None, None, e.get("facts", ""),
                       e.get("level", ""), e.get("risk", "")])
        style_table(ws, [8, 14, 36, 48, 52, 40, 9, 9, 10, 8, 40])
        for n in range(2, ws.max_row + 1):
            fill = PatternFill("solid", fgColor=PALETTE[(n - 2) % len(PALETTE)])
            for col in range(1, len(exc_headers) + 1):
                ws.cell(n, col).fill = fill
            lvl = str(ws.cell(n, 10).value or "")
            for mark, color in LEVEL_FILL.items():
                if lvl.startswith(mark):
                    ws.cell(n, 10).fill = PatternFill("solid", fgColor=color)
            for col in (3, 4):
                ws.cell(n, col).font = Font(name=FONT, bold=True, size=10, color="C00000")
        ws.insert_rows(1)  # form/source line above the table; LEN formulas written after the shift
        for n in range(3, ws.max_row + 1):
            ws.cell(n, 7).value = f"=LEN(C{n})"
            ws.cell(n, 8).value = f"=LEN(D{n})"
        c = ws.cell(1, 1, value=f"剂型：{exc.get('form', '')}｜出处：{exc.get('source', '')}")
        c.font = Font(name=FONT, bold=True, size=11)
        ws.freeze_panes = "C3"
        ws.auto_filter.ref = f"A2:K{ws.max_row}"
        ws.row_dimensions[2].height = 30

    # --- 设计修改清单 ---
    ws = wb.create_sheet("设计修改清单")
    ws.append(["#", "优先级", "位置", "问题", "建议处理", "状态"])
    for i, it in enumerate(data.get("issues", []), 1):
        ws.append([i, it.get("priority", "建议"), it.get("location", ""), it.get("problem", ""),
                   it.get("fix", ""), it.get("status", "待处理")])
    style_table(ws, [5, 8, 24, 50, 50, 10])
    for n in range(2, ws.max_row + 1):
        p = ws.cell(n, 2).value
        if p in PRI_FILL:
            ws.cell(n, 2).fill = PatternFill("solid", fgColor=PRI_FILL[p])

    # --- 事实依据 ---
    ws = wb.create_sheet("事实依据")
    ws.append(["编号", "事实", "出处", "备注"])
    for f in data.get("facts", []):
        ws.append([f.get("id", ""), f.get("fact", ""), f.get("source", ""), f.get("note", "")])
    style_table(ws, [6, 56, 50, 48])

    wb.save(out)
    print(f"saved {out}  rows={len(data['rows'])} marketing={len(data.get('marketing') or [])} "
          f"explainers={len(data.get('explainers') or [])} "
          f"excipients={len((data.get('excipients') or {}).get('items') or [])} "
          f"issues={len(data.get('issues', []))} facts={len(data.get('facts', []))}")


if __name__ == "__main__":
    if sys.argv[1:] == ["--stamp-only"]:
        print(stamp())
        sys.exit(0)
    if len(sys.argv) != 3:
        sys.exit("usage: build_xlsx.py copy.json output.xlsx | build_xlsx.py --stamp-only")
    main(sys.argv[1], sys.argv[2])
