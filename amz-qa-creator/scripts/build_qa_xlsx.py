# -*- coding: utf-8 -*-
"""合并 QA 批次 -> 全量质检 -> 导出 xlsx/csv/txt。
用法: python build_qa_xlsx.py <批次目录> <facts.json路径> <输出前缀(不含扩展名)>
批次目录内含 qa_batchN.json；若有 manifest.json 则按其 status=done 列表加载（断点续跑）。
facts.json 示例: {"brand":"ZAB","dose_mg":1000,"unit_count":90,"servings":45,"price":49.99,
                  "coupon":5.0,"dose_softgels":2,"form":"softgel","flavor":"unflavored","age_range":"adult"}
依赖: openpyxl"""
import csv, glob, json, os, re, sys, io, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

if len(sys.argv) < 4:
    sys.exit("usage: python build_qa_xlsx.py <batches_dir> <facts.json> <out_prefix>")
DIR, FACTS_PATH, PREFIX = sys.argv[1], sys.argv[2], sys.argv[3]
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BANNED_PATH = os.path.join(SCRIPT_DIR, "..", "references", "banned_words.txt")

# ---------- 加载批次（manifest 优先，支持断点续跑） ----------
def load_batch(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)

items, manifest = [], {"batches": [], "updated": datetime.datetime.now().isoformat(timespec="seconds")}
mpath = os.path.join(DIR, "manifest.json")
if os.path.exists(mpath):
    old = json.load(open(mpath, encoding="utf-8"))
    files = [os.path.join(DIR, b["file"]) for b in old.get("batches", []) if b.get("status") == "done"]
    if not files:
        files = sorted(glob.glob(os.path.join(DIR, "qa_batch*.json")))
else:
    files = sorted(glob.glob(os.path.join(DIR, "qa_batch*.json")))
if not files:
    sys.exit("no qa_batch*.json found in " + DIR)
for fp in files:
    data = load_batch(fp)
    items.extend(data)
    manifest["batches"].append({"file": os.path.basename(fp), "count": len(data), "status": "done"})
with open(mpath, "w", encoding="utf-8") as f:
    json.dump(manifest, f, ensure_ascii=False, indent=1)

with open(FACTS_PATH, encoding="utf-8") as f:
    facts = json.load(f)

# ---------- 敏感词清单 ----------
def load_banned(path):
    red, warn, cur = [], [], None
    with open(path, encoding="utf-8") as f:
        for ln in f:
            ln = ln.strip()
            if not ln or ln.startswith("#"):
                continue
            if ln == "[REDLINE]":
                cur = red
            elif ln == "[WARNING]":
                cur = warn
            elif cur is not None:
                cur.append(ln.lower())
    return red, warn

REDLINE, WARNWORDS = load_banned(BANNED_PATH)

def hit_words(text, words):
    t = text.lower()
    return [w for w in words if re.search(r"(?<![a-z])" + re.escape(w) + r"(?![a-z])", t)]

# ---------- 质检 ----------
qc, fails, warns = [], [], []
n = len(items)
target = manifest.get("target", 200)
qc.append(("总数", f"{n} (目标{target}/最低100)", "PASS" if n >= 100 else "FAIL"))

type_cnt = collections.Counter(t for r in items for t in r["t"])
src_cnt = collections.Counter(r["s"] for r in items)
act_cnt = collections.Counter(r["act"] for r in items)
qc.append(("类型分布(可多标)", str(dict(type_cnt)), "INFO"))
paa_pct = src_cnt.get("google_paa", 0) / n * 100
qc.append(("google_paa占比", f"{paa_pct:.1f}% (上限20%)", "PASS" if paa_pct <= 20 else "FAIL"))
qc.append(("建议动作分布", str(dict(act_cnt)), "INFO"))

BRANDS = ["mitopure", "timeline", "codeage", "double wood", "jarrow", "life extension", "neurogan",
          "omre", "donotage", "california gold", "pure encapsulation", "renue", "onset", "aeternum",
          "totaria", "dncdtyy", "bizyac", "mityvac", "cystore"]
brand_cnt = collections.Counter()
for r in items:
    text = (r["q"] + " " + r["a"]).lower()
    for b in BRANDS:
        if b in text:
            brand_cnt[b] += 1
cap = n * 0.15
qc.append(("品牌词频次", f"{dict(brand_cnt)} (上限{cap:.0f}/品牌)", "PASS" if all(v <= cap for v in brand_cnt.values()) else "FAIL"))
prev, consec = None, True
for r in items:
    text = (r["q"] + " " + r["a"]).lower()
    cur = next((b for b in BRANDS if b in text), None)
    if cur and cur == prev:
        consec = False
    prev = cur
qc.append(("同品牌连续出现", "", "PASS" if consec else "FAIL"))

FORM = ["gummies", "gummy", "powder", "liquid drops"]
fm = [(i, k) for i, r in enumerate(items, 1) for k in r["k"] if any(f in k.lower() for f in FORM)]
qc.append(("形态错配词混入埋词列", str(fm) if fm else "0处", "PASS" if not fm else "FAIL"))

qs = [r["q"].strip().lower() for r in items]
dup = [q for q, c in collections.Counter(qs).items() if c > 1]
qc.append(("问句完全重复", str(dup) if dup else "无", "PASS" if not dup else "FAIL"))

# 问句开场句式分布
openers = collections.Counter(q.split()[0] if q.split() else "" for q in qs)
top_openers = {w: f"{c}({c/n*100:.0f}%)" for w, c in openers.most_common(8)}
op_fail = {w: c for w, c in openers.items() if c / n > 0.30}
qc.append(("问句开场句式分布", str(top_openers), "PASS" if not op_fail else f"FAIL(>30%: {op_fail})"))

# 问句前三词 n-gram 重复
tri = collections.Counter(" ".join(re.sub(r"[^a-z0-9 ]", "", q).split()[:3]) for q in qs)
tri_rep = {k: v for k, v in tri.items() if v > max(3, n * 0.03) and len(k.split()) == 3}
qc.append(("问句前三词重复", str(tri_rep) if tri_rep else "无异常", "PASS" if not tri_rep else "WARNING"))
if tri_rep:
    warns.append(f"问句前三词高频重复: {tri_rep}")

# 答案开头模式：全局高频 + 连续3条
ans_open = collections.Counter(" ".join(re.sub(r"[^a-z0-9' ]", "", r["a"].lower()).split()[:3]) for r in items)
ans_hot = {k: v for k, v in ans_open.most_common(5) if v > 5}
run, run_hits = 1, []
for i in range(1, n):
    same = " ".join(re.sub(r"[^a-z0-9' ]", "", items[i]["a"].lower()).split()[:3]) == \
           " ".join(re.sub(r"[^a-z0-9' ]", "", items[i - 1]["a"].lower()).split()[:3])
    run = run + 1 if same else 1
    if run == 3:
        run_hits.append(i + 1)
qc.append(("答案开头模式", f"高频:{ans_hot} 连续3条行:{run_hits}" if (ans_hot or run_hits) else "无异常",
           "PASS" if not run_hits else "WARNING"))
if ans_hot or run_hits:
    warns.append(f"答案开头模式需人工润色: 高频={ans_hot} 连续={run_hits}")

# 敏感词扫描（答句红线=FAIL；问句红线=WARNING，买家自然提问可保留；医疗敏感词=人工复核）
red_a = [(i, hit_words(r["a"], REDLINE)) for i, r in enumerate(items, 1)]
red_a = [(i, w) for i, w in red_a if w]
red_q = [(i, hit_words(r["q"], REDLINE)) for i, r in enumerate(items, 1)]
red_q = [(i, w) for i, w in red_q if w]
warn_hits = [(i, hit_words(r["q"] + " " + r["a"], WARNWORDS)) for i, r in enumerate(items, 1)]
warn_hits = [(i, w) for i, w in warn_hits if w]
qc.append(("红线词-答句", str(red_a) if red_a else "0处", "PASS" if not red_a else "FAIL"))
qc.append(("红线词-问句(买家自然提问)", str(red_q) if red_q else "0处", "PASS" if not red_q else "WARNING"))
qc.append(("医疗敏感词(安全问答专用)", f"{len(warn_hits)}行命中，需人工确认为'请咨询医生'语境: {[i for i, _ in warn_hits]}" if warn_hits else "0处", "INFO" if not warn_hits else "WARNING"))
if red_a:
    fails.append(f"答句红线词命中行: {red_a}")
if red_q:
    warns.append(f"问句红线词(确认买家自然语境): {red_q}")
if warn_hits:
    warns.append(f"医疗敏感词待复核行: {[i for i, _ in warn_hits]}")

# 事实锚点比对（规格类=FAIL 防幻觉；价格/剂量=WARNING，允许竞品对比语境）
dose = facts.get("dose_mg")
allow = {
    "price": {facts.get("price"), facts.get("coupon")},
    "softgels": {facts.get("unit_count"), facts.get("dose_softgels", 2), 1},
    "count": {facts.get("unit_count")},
    "servings": {facts.get("servings"), 1},
    "mg": {dose, dose / 2 if dose else None, dose * 2 if dose else None},
}
allow = {k: {x for x in v if x is not None} for k, v in allow.items()}
pats = [(r"\$(\d+(?:\.\d+)?)", "price", 1), (r"(\d+(?:\.\d+)?)\s?mg\b", "mg", 1),
        (r"(\d+)\s?softgels?", "softgels", 1), (r"(\d+)\s?count", "count", 1),
        (r"(\d+)\s?servings?", "servings", 1)]
STRICT_KEYS = {"softgels", "count", "servings"}
fact_bad, fact_warn = [], []
for i, r in enumerate(items, 1):
    for pat, key, _ in pats:
        for m in re.finditer(pat, r["a"].lower()):
            val = float(m.group(1))
            if allow[key] and val not in allow[key]:
                (fact_bad if key in STRICT_KEYS else fact_warn).append((i, key, val, sorted(allow[key])))
qc.append(("事实锚点-规格(软胶囊数/份数)", str(fact_bad) if fact_bad else "全部一致", "PASS" if not fact_bad else "FAIL"))
qc.append(("事实锚点-价格/剂量(允许竞品对比)", str(fact_warn) if fact_warn else "全部一致", "PASS" if not fact_warn else "WARNING"))
if fact_bad:
    fails.append(f"规格事实不一致行: {fact_bad}")
if fact_warn:
    warns.append(f"价格/剂量对比语境待确认行: {fact_warn}")

# 核心参数覆盖
core_params = {"剂量": str(facts.get("dose_mg", "")) + "mg" if facts.get("dose_mg") else None,
               "规格": str(facts.get("unit_count", "")) if facts.get("unit_count") else None,
               "形态": facts.get("form"), "口味": facts.get("flavor"), "人群": facts.get("age_range")}
cover = {}
for name, p in core_params.items():
    if not p:
        continue
    cover[name] = any(p.lower() in (r["q"] + " " + r["a"]).lower() for r in items)
qc.append(("核心参数覆盖", str(cover), "PASS" if all(cover.values()) else "FAIL"))

overall = "FAIL" if fails else ("PASS(含WARNING待人工复核)" if warns else "PASS(全部通过)")
qc.append(("总判定", overall, "INFO"))

# ---------- 输出 ----------
rows = [[i, r["q"], r["a"], "/".join(r["t"]), ", ".join(r["k"]), r["d"], r["s"], r["act"]] for i, r in enumerate(items, 1)]
HEADERS = ["序号", "Question", "Answer", "QA类型", "埋入关键词", "覆盖维度", "来源标签", "建议动作"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "QA总表"
ws.append(HEADERS)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E79")
    c.alignment = Alignment(vertical="center")
for row in rows:
    ws.append(row)
    if row[7] == "建议人工微调":
        ws.cell(row=row[0] + 1, column=8).fill = PatternFill("solid", fgColor="FFF2CC")
for i, w in enumerate([6, 42, 90, 16, 34, 16, 14, 14], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("统计与质检")
ws2.append(["AMZ-QA-Creator 质检报告", datetime.date.today().isoformat()])
ws2["A1"].font = Font(bold=True, size=13)
ws2.append(["检查项", "结果", "判定"])
for c in ws2[2]:
    c.font = Font(bold=True)
for name, res, verdict in qc:
    ws2.append([name, res, verdict])
    if verdict.startswith("FAIL"):
        ws2.cell(row=ws2.max_row, column=3).fill = PatternFill("solid", fgColor="F8CBAD")
    elif verdict == "WARNING":
        ws2.cell(row=ws2.max_row, column=3).fill = PatternFill("solid", fgColor="FFF2CC")
ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 110
ws2.column_dimensions["C"].width = 14

xlsx_path = PREFIX + ".xlsx"
wb.save(xlsx_path)
with open(PREFIX + ".csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(HEADERS)
    w.writerows(rows)
with open(PREFIX + ".txt", "w", encoding="utf-8") as f:
    for row in rows:
        f.write(f"{row[0]}. Q: {row[1]}\n   A: {row[2]}\n\n")

print(f"saved -> {xlsx_path} / .csv / .txt | items={n}")
for name, res, verdict in qc:
    print(f"[{verdict}] {name}: {res[:160]}")
sys.exit(1 if fails else 0)
