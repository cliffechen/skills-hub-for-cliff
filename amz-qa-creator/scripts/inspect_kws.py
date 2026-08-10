# -*- coding: utf-8 -*-
"""解析词库六sheet xlsx -> JSON。
用法: python inspect_kws.py <词库xlsx路径> <输出json路径>
依赖: openpyxl"""
import json, sys, io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

if len(sys.argv) < 3:
    sys.exit("usage: python inspect_kws.py <kws.xlsx> <out.json>")
PATH, OUT = sys.argv[1], sys.argv[2]

wb = openpyxl.load_workbook(PATH, data_only=True)
print("=== SHEETS ===")
for ws in wb.worksheets:
    headers = [str(c.value) if c.value is not None else "" for c in ws[1]]
    print(f"[{ws.title}] rows={ws.max_row} cols={ws.max_column} | " + " | ".join(headers))

def dump(ws, max_rows=500):
    headers = [str(c.value) if c.value is not None else f"col{i}" for i, c in enumerate(ws[1])]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=min(ws.max_row, max_rows + 1), values_only=True):
        if all(v is None for v in r):
            continue
        rows.append({headers[i]: (str(v) if v is not None else "") for i, v in enumerate(r)})
    return rows

out = {name: dump(wb[name]) for name in wb.sheetnames}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print(f"dumped -> {OUT} | sheets: {len(out)}")
